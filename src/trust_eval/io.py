from __future__ import annotations

import csv
import json
import re
from pathlib import Path
from typing import Any, Iterable


ALIASES = {
    "case_id": {"case_id", "id", "编号", "序号", "题号"},
    "query": {"query", "question", "prompt", "用户问题", "问题", "输入"},
    "answer": {"answer", "response", "model_answer", "模型回答", "回复", "输出"},
    "conversation": {"conversation", "history", "对话历史", "上下文"},
    "expected_behavior": {"expected_behavior", "expectation", "ground_truth", "参考标准", "期望行为", "标准答案"},
    "exact_answer": {"exact_answer", "gold_answer", "唯一答案", "精确答案"},
    "required_terms": {"required_terms", "must_include", "必含词", "关键点"},
    "forbidden_terms": {"forbidden_terms", "must_not_include", "禁用词", "禁止内容"},
    "category": {"category", "type", "题型", "分类"},
    "attachment_state": {"attachment_state", "附件状态", "图片状态"},
    "action_state": {"action_state", "执行状态", "操作状态"},
    "available_tools": {"available_tools", "可用工具"},
    "tool_trace": {"tool_trace", "工具轨迹", "调用轨迹"},
}


def _canonical_key(value: Any) -> str:
    raw = str(value or "").strip().lower()
    for key, aliases in ALIASES.items():
        if raw in {alias.lower() for alias in aliases}:
            return key
    return raw


def _as_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return str(value).strip()


def _split_terms(value: Any) -> list[str]:
    if value is None or value == "":
        return []
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    text = str(value).strip()
    if text.startswith("["):
        try:
            parsed = json.loads(text)
            if isinstance(parsed, list):
                return [str(item).strip() for item in parsed if str(item).strip()]
        except json.JSONDecodeError:
            pass
    return [item.strip() for item in re.split(r"[;,；，\n]", text) if item.strip()]


def _parse_jsonish(value: Any, fallback: Any) -> Any:
    if isinstance(value, (list, dict)):
        return value
    if not value:
        return fallback
    try:
        return json.loads(str(value))
    except json.JSONDecodeError:
        return fallback


def normalize_records(records: Iterable[dict[str, Any]], source: str) -> list[dict[str, Any]]:
    cases: list[dict[str, Any]] = []
    for index, raw in enumerate(records, start=1):
        nested_truth = raw.get("ground_truth") if isinstance(raw.get("ground_truth"), dict) else {}
        nested_context = raw.get("context") if isinstance(raw.get("context"), dict) else {}
        nested_rubric = raw.get("rubric") if isinstance(raw.get("rubric"), dict) else {}
        nested_review = raw.get("review") if isinstance(raw.get("review"), dict) else {}
        row = {_canonical_key(key): value for key, value in raw.items()}
        query = _as_text(row.get("query"))
        answer = _as_text(row.get("answer"))
        if not query and not answer:
            continue
        case_id = _as_text(row.get("case_id")) or f"case-{index:03d}"
        extra = {
            str(key): value
            for key, value in row.items()
            if key not in set(ALIASES) | {"context", "rubric", "review", "schema_version"} and value not in (None, "")
        }
        cases.append(
            {
                "schema_version": "1.0",
                "case_id": case_id,
                "query": query,
                "answer": answer,
                "conversation": _parse_jsonish(row.get("conversation"), _as_text(row.get("conversation"))),
                "ground_truth": {
                    "expected_behavior": _as_text(nested_truth.get("expected_behavior", row.get("expected_behavior"))),
                    "exact_answer": _as_text(nested_truth.get("exact_answer", row.get("exact_answer"))) or None,
                    "required_terms": _split_terms(nested_truth.get("required_terms", row.get("required_terms"))),
                    "forbidden_terms": _split_terms(nested_truth.get("forbidden_terms", row.get("forbidden_terms"))),
                },
                "context": {
                    "evaluation_time": nested_context.get("evaluation_time"),
                    "timezone": nested_context.get("timezone"),
                    "attachment_state": _as_text(nested_context.get("attachment_state", row.get("attachment_state"))) or "unknown",
                    "available_tools": _split_terms(nested_context.get("available_tools", row.get("available_tools"))),
                    "tool_trace": _parse_jsonish(nested_context.get("tool_trace", row.get("tool_trace")), []),
                    "action_state": _as_text(nested_context.get("action_state", row.get("action_state"))) or "unknown",
                    "source": source,
                    "extra": nested_context.get("extra", extra),
                },
                "rubric": {**nested_rubric, "category": _as_text(nested_rubric.get("category", row.get("category")))},
                "review": {
                    "status": nested_review.get("status", "draft"),
                    "reviewer": nested_review.get("reviewer", ""),
                    "notes": nested_review.get("notes", ""),
                },
            }
        )
    if not cases:
        raise ValueError("没有找到包含 query/question 与 answer/response 的有效记录")
    return cases


def _records_from_rows(rows: list[list[Any]]) -> list[dict[str, Any]]:
    rows = [row for row in rows if any(cell not in (None, "") for cell in row)]
    if not rows:
        return []
    headers = [_as_text(cell) or f"column_{index + 1}" for index, cell in enumerate(rows[0])]
    return [dict(zip(headers, row)) for row in rows[1:]]


def _read_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _read_excel(path: Path) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    records: list[dict[str, Any]] = []
    for sheet in workbook.worksheets:
        rows = [list(row) for row in sheet.iter_rows(values_only=True)]
        for record in _records_from_rows(rows):
            record.setdefault("source_sheet", sheet.title)
            records.append(record)
    return records


def _read_json(path: Path) -> list[dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(payload, list):
        return payload
    if isinstance(payload, dict):
        for key in ("cases", "items", "data", "records"):
            if isinstance(payload.get(key), list):
                return payload[key]
        return [payload]
    raise ValueError("JSON 顶层必须是对象或数组")


def _read_jsonl(path: Path) -> list[dict[str, Any]]:
    records = []
    for line_number, line in enumerate(path.read_text(encoding="utf-8").splitlines(), start=1):
        if not line.strip():
            continue
        try:
            records.append(json.loads(line))
        except json.JSONDecodeError as exc:
            raise ValueError(f"JSONL 第 {line_number} 行无法解析: {exc}") from exc
    return records


def _markdown_table(text: str) -> list[dict[str, Any]]:
    lines = [line.strip() for line in text.splitlines() if line.strip().startswith("|")]
    if len(lines) < 2:
        return []
    rows = [[cell.strip() for cell in line.strip("|").split("|")] for line in lines]
    if all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in rows[1]):
        rows.pop(1)
    return _records_from_rows(rows)


def _labeled_blocks(text: str) -> list[dict[str, Any]]:
    records: list[dict[str, Any]] = []
    current: dict[str, Any] = {}
    label_pattern = re.compile(
        r"^(case_id|id|编号|题号|query|question|问题|用户问题|answer|response|回答|模型回答|ground_truth|期望行为|参考标准)\s*[:：]\s*(.*)$",
        re.I,
    )
    for line in text.splitlines() + ["---"]:
        stripped = line.strip()
        if stripped in {"---", "***"}:
            if current:
                records.append(current)
                current = {}
            continue
        match = label_pattern.match(stripped)
        if match:
            current[match.group(1)] = match.group(2)
    return records


def _read_text(path: Path) -> list[dict[str, Any]]:
    text = path.read_text(encoding="utf-8")
    return _markdown_table(text) or _labeled_blocks(text)


def _read_docx(path: Path) -> list[dict[str, Any]]:
    from docx import Document

    document = Document(path)
    records: list[dict[str, Any]] = []
    for table in document.tables:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        records.extend(_records_from_rows(rows))
    if records:
        return records
    return _labeled_blocks("\n".join(paragraph.text for paragraph in document.paragraphs))


def read_cases(path: str | Path) -> list[dict[str, Any]]:
    source = Path(path).expanduser().resolve()
    if not source.exists():
        raise FileNotFoundError(source)
    suffix = source.suffix.lower()
    readers = {
        ".csv": _read_csv,
        ".xlsx": _read_excel,
        ".xlsm": _read_excel,
        ".json": _read_json,
        ".jsonl": _read_jsonl,
        ".md": _read_text,
        ".txt": _read_text,
        ".docx": _read_docx,
    }
    if suffix not in readers:
        raise ValueError(f"暂不支持 {suffix or '无扩展名'}；支持 CSV/XLSX/JSON/JSONL/MD/TXT/DOCX")
    return normalize_records(readers[suffix](source), str(source))


def read_jsonl(path: str | Path) -> list[dict[str, Any]]:
    return _read_jsonl(Path(path))


def write_jsonl(path: str | Path, items: Iterable[dict[str, Any]]) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(
        "".join(json.dumps(item, ensure_ascii=False) + "\n" for item in items),
        encoding="utf-8",
    )


def write_json(path: str | Path, payload: Any) -> None:
    target = Path(path)
    target.parent.mkdir(parents=True, exist_ok=True)
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
